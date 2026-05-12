"""pytest 夹具 — 创建测试用 .docx 文件和共享资源。"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Generator

import docx
import pytest


FIXTURES_DIR = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="session")
def fixtures_dir() -> Path:
    """测试夹具文件目录。"""
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    return FIXTURES_DIR


@pytest.fixture(scope="session")
def simple_docx_path(fixtures_dir: Path) -> Path:
    """创建一个简单的测试用 .docx 文件。"""
    filepath = fixtures_dir / "simple.docx"

    if filepath.exists():
        return filepath

    doc = docx.Document()

    # 添加标题
    doc.add_heading("第一章 引言", level=1)

    doc.add_paragraph("这是正文段落，介绍了文档的背景和目的。")
    doc.add_paragraph("本段落是第二个正文段落。")

    doc.add_heading("1.1 研究背景", level=2)
    doc.add_paragraph("在研究背景中，我们讨论了当前领域的发展状况。")

    doc.add_heading("1.2 相关工作", level=2)
    doc.add_paragraph("相关工作部分回顾了已有文献。")

    doc.add_heading("第二章 方法", level=1)
    doc.add_paragraph("本章介绍了本文提出的方法。")

    # 添加列表
    for item in ["步骤一：数据预处理", "步骤二：特征提取", "步骤三：模型训练"]:
        p = doc.add_paragraph(item, style="List Bullet")
    p = doc.add_paragraph(item, style="List Number")

    # 添加表格
    doc.add_heading("2.1 实验数据", level=2)
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "方法"
    table.cell(0, 1).text = "准确率"
    table.cell(0, 2).text = "F1 值"
    table.cell(1, 0).text = "方法A"
    table.cell(1, 1).text = "90.5%"
    table.cell(1, 2).text = "0.89"
    table.cell(2, 0).text = "方法B"
    table.cell(2, 1).text = "92.3%"
    table.cell(2, 2).text = "0.91"

    doc.save(str(filepath))
    return filepath


@pytest.fixture(scope="session")
def simple_docx_bytes(simple_docx_path: Path) -> bytes:
    """测试用 .docx 文件的字节流。"""
    return simple_docx_path.read_bytes()


@pytest.fixture
def docx_document(simple_docx_path: Path) -> docx.Document:
    """返回一个 python-docx Document 对象。"""
    return docx.Document(str(simple_docx_path))


@pytest.fixture
def default_ruleset_path() -> Path:
    """默认规则集路径。"""
    return Path(__file__).parent.parent / "rulesets" / "default.yaml"


@pytest.fixture
def academic_ruleset_path() -> Path:
    """学术规则集路径。"""
    return Path(__file__).parent.parent / "rulesets" / "academic.yaml"


@pytest.fixture
def technical_ruleset_path() -> Path:
    """技术文档规则集路径。"""
    return Path(__file__).parent.parent / "rulesets" / "technical.yaml"
